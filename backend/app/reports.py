"""
Admin Excel report (FR11): one row per conversation.
Rows are only ever APPENDED; the file is created if it does not exist.
"""
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

REPORT_PATH = Path(__file__).resolve().parent.parent / "reports" / "conversations.xlsx"
SHEET = "Conversations"

# (column name, width) - the order here is the order in Excel
COLUMNS = [
    ("Conversation ID", 16), ("Date", 12), ("Time", 9),
    ("Patient Name", 22), ("Phone", 14), ("Topics", 32),
    ("Booking Status", 15), ("Doctor", 24), ("Speciality", 18),
    ("Appointment Date", 17), ("Appointment Time", 17), ("Invoice No", 12),
    ("Summary", 70),
]
BOOKED_FILL = PatternFill("solid", fgColor="D1FAE5")      # light green
NOT_BOOKED_FILL = PatternFill("solid", fgColor="F3F4F6")  # light grey


def _create_workbook() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    ws.append([name for name, _ in COLUMNS])
    for col, (name, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F766E")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[cell.column_letter].width = width
    ws.freeze_panes = "A2"          # header stays visible when scrolling
    return wb


def append_conversation_row(row: dict) -> dict:
    """row keys = column names. Missing keys become empty cells."""
    REPORT_PATH.parent.mkdir(parents=True, exist_ok=True)
    try:
        wb = load_workbook(REPORT_PATH) if REPORT_PATH.exists() else _create_workbook()
        ws = wb[SHEET]
        ws.append([row.get(name, "") for name, _ in COLUMNS])

        r = ws.max_row
        status_col = [n for n, _ in COLUMNS].index("Booking Status") + 1
        status_cell = ws.cell(row=r, column=status_col)
        status_cell.fill = BOOKED_FILL if status_cell.value == "Booked" else NOT_BOOKED_FILL
        ws.cell(row=r, column=len(COLUMNS)).alignment = Alignment(wrap_text=True, vertical="top")
        ws.auto_filter.ref = f"A1:{ws.cell(row=r, column=len(COLUMNS)).coordinate}"

        wb.save(REPORT_PATH)
    except PermissionError:
        return {"ok": False, "error": "The Excel report is open in another program. "
                                      "Close it and try again."}
    return {"ok": True, "row": r, "path": str(REPORT_PATH)}