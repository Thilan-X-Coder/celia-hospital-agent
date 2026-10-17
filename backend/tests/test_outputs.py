from openpyxl import load_workbook

from app import invoice, reports, tools
from app.conversation_end import finish_conversation
from app.sessions import Session

PATIENT = dict(full_name="Jane Smith", phone="0712345678",
               date_of_birth="1990-02-15", reason="Skin rash & itching")


def fake_summary(session):           # no LLM in tests: fast, free, offline
    return "Test summary."


def _book():
    slot_id = tools.get_available_slots(3)["slots"][0]["slot_id"]
    ref = tools.prepare_booking(slot_id, **PATIENT)["booking_ref"]
    return tools.confirm_booking(ref)


def _session_with_message(text="What are your visiting hours?"):
    s = Session()
    s.messages.append({"role": "user", "content": text})
    return s


def _sheet():
    return load_workbook(reports.REPORT_PATH)["Conversations"]


# ---------- invoice PDF (FR9) ----------
def test_booking_creates_invoice_pdf(test_db):
    result = _book()
    assert result["invoice_ready"] is True
    pdf = invoice.INVOICES_DIR / "INV-0001_JaneSmith.pdf"
    assert pdf.exists()
    assert pdf.read_bytes()[:4] == b"%PDF"       # every PDF file starts with %PDF


# ---------- Excel report (FR10, FR11) ----------
def test_excel_is_created_with_header_and_row(test_db):
    s = _session_with_message()
    s.tools_used.append("get_hospital_info")
    assert finish_conversation(s, summarizer=fake_summary)["ok"] is True

    ws = _sheet()
    assert ws.cell(row=1, column=1).value == "Conversation ID"
    assert ws.cell(row=2, column=6).value == "Hospital info"      # Topics
    assert ws.cell(row=2, column=7).value == "Not booked"         # Booking Status


def test_rows_are_appended_not_overwritten(test_db):
    for _ in range(3):
        finish_conversation(_session_with_message(), summarizer=fake_summary)
    assert _sheet().max_row == 4                                  # header + 3 rows


def test_booked_conversation_row(test_db):
    s = _session_with_message("I want to see a dermatologist")
    s.booking = _book()
    s.patient = {"full_name": "Jane Smith", "phone": "0712345678"}
    finish_conversation(s, summarizer=fake_summary)

    row = [c.value for c in _sheet()[2]]
    assert "Booked" in row
    assert "INV-0001" in row
    assert "Jane Smith" in row


def test_conversation_is_saved_only_once(test_db):
    s = _session_with_message()
    finish_conversation(s, summarizer=fake_summary)
    finish_conversation(s, summarizer=fake_summary)
    assert _sheet().max_row == 2


def test_row_is_saved_even_when_llm_fails(test_db):
    def broken_summary(session):
        raise RuntimeError("LLM is down")

    result = finish_conversation(_session_with_message(), summarizer=broken_summary)
    assert result["ok"] is True
    assert "[auto summary]" in result["summary"]