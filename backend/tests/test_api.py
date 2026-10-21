"""
HTTP-wiring tests for app/main.py.

These do NOT call the real LLM: app.main.run_turn is monkeypatched per test, and
app.main.finish_conversation is monkeypatched (where a summary matters) to a wrapper
that passes a fake summarizer straight into the real conversation_end.finish_conversation
- patching app.summary.make_summary would NOT work, because finish_conversation's
default argument (summarizer=make_summary) is bound once at import time.
"""
from datetime import datetime, timedelta

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app import conversation_end, main, reports, sessions, tools


def _fake_finish(monkeypatch):
    """Route app.main's finish_conversation calls to a fake summarizer, no real LLM."""
    monkeypatch.setattr(
        main,
        "finish_conversation",
        lambda session: conversation_end.finish_conversation(
            session, summarizer=lambda s: "A fake summary."
        ),
    )


@pytest.fixture
def client(test_db):
    return TestClient(main.app)


def _start(client):
    r = client.post("/api/chat/start")
    assert r.status_code == 200
    body = r.json()
    assert body["session_id"].startswith("CONV-")
    assert body["greeting"]
    return body["session_id"]


# ---------------- happy path: a) plain turn, no booking ----------------
def test_normal_turn_has_no_booking_or_invoice(client, monkeypatch):
    monkeypatch.setattr(main, "run_turn", lambda session, text, debug=False: "a normal reply")

    session_id = _start(client)
    r = client.post("/api/chat", json={"session_id": session_id, "message": "hello"})

    assert r.status_code == 200
    body = r.json()
    assert body["reply"] == "a normal reply"
    assert body["booking"] is None
    assert body["invoice_url"] is None
    assert body["conversation_ended"] is False
    assert body["pending_booking"] is None


# ---------------- pending_booking appears on the turn prepare_booking succeeds ----------------
def test_chat_includes_pending_booking_after_prepare(client, monkeypatch):
    doctor = tools.list_doctors()["doctors"][0]
    slot = tools.get_available_slots(doctor["doctor_id"])["slots"][0]

    calls = {"n": 0}

    def fake_run_turn(session, text, debug=False):
        session.turn_no += 1  # real run_turn does this as its very first step
        calls["n"] += 1
        if calls["n"] == 1:
            prep = tools.prepare_booking(
                slot["slot_id"], "Jane Smith", "0712345678", "1990-05-15", "General checkup",
            )
            assert prep["ok"], prep
            session.pending_refs[prep["booking_ref"]] = session.turn_no
            return "Shall I confirm this booking?"
        return "Anything else I can help with?"

    monkeypatch.setattr(main, "run_turn", fake_run_turn)

    session_id = _start(client)
    r = client.post("/api/chat", json={"session_id": session_id, "message": "book it"})

    assert r.status_code == 200
    body = r.json()
    assert body["booking"] is None

    pb = body["pending_booking"]
    assert pb is not None
    assert pb["patient_name"] == "Jane Smith"
    assert pb["doctor"] == doctor["full_name"]
    assert pb["speciality"] == doctor["speciality"]
    assert pb["date"] == slot["slot_date"]
    assert pb["time"] == slot["start_time"]
    assert pb["fee_lkr"] == doctor["consultation_fee"]

    # booking_ref must never reach the frontend
    assert "booking_ref" not in r.text

    # the next turn (no new prepare_booking call) must not repeat the pending booking
    r2 = client.post("/api/chat", json={"session_id": session_id, "message": "thanks"})
    assert r2.json()["pending_booking"] is None


# ---------------- b) + c): a real booking appears once, then the PDF is fetchable ----------------
def test_booking_and_invoice_appear_once_then_pdf_is_fetchable(client, monkeypatch):
    doctor = tools.list_doctors()["doctors"][0]
    slot = tools.get_available_slots(doctor["doctor_id"])["slots"][0]

    calls = {"n": 0}

    def fake_run_turn(session, text, debug=False):
        calls["n"] += 1
        if calls["n"] == 1:
            prep = tools.prepare_booking(
                slot["slot_id"], "Jane Smith", "0712345678", "1990-05-15", "General checkup",
            )
            assert prep["ok"], prep
            session.booking = tools.confirm_booking(prep["booking_ref"])
            assert session.booking["ok"], session.booking
            return "You're all booked."
        return "Anything else I can help with?"

    monkeypatch.setattr(main, "run_turn", fake_run_turn)

    session_id = _start(client)

    r1 = client.post("/api/chat", json={"session_id": session_id, "message": "book it"})
    assert r1.status_code == 200
    body1 = r1.json()
    assert body1["booking"] is not None
    assert body1["booking"]["doctor"] == doctor["full_name"]
    assert body1["booking"]["speciality"] == doctor["speciality"]
    assert body1["invoice_url"] is not None
    assert body1["invoice_url"].startswith("/invoices/") and body1["invoice_url"].endswith(".pdf")

    # the same booking must NOT be sent again on the next turn
    r2 = client.post("/api/chat", json={"session_id": session_id, "message": "thanks"})
    assert r2.status_code == 200
    body2 = r2.json()
    assert body2["booking"] is None
    assert body2["invoice_url"] is None

    r3 = client.get(body1["invoice_url"])
    assert r3.status_code == 200
    assert r3.content[:4] == b"%PDF"


# ---------------- d): goodbye triggers an automatic save ----------------
def test_goodbye_saves_conversation_automatically(client, monkeypatch):
    def fake_run_turn(session, text, debug=False):
        session.messages.append({"role": "user", "content": text})
        return "Goodbye, take care!"

    def fake_finish(session):
        return conversation_end.finish_conversation(session, summarizer=lambda s: "A fake summary.")

    monkeypatch.setattr(main, "run_turn", fake_run_turn)
    monkeypatch.setattr(main, "finish_conversation", fake_finish)

    session_id = _start(client)
    r = client.post("/api/chat", json={"session_id": session_id, "message": "bye, thank you"})

    assert r.status_code == 200
    assert r.json()["conversation_ended"] is True

    wb = openpyxl.load_workbook(reports.REPORT_PATH)
    ws = wb[reports.SHEET]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert any(row[0] == session_id for row in rows)


# ---------------- e): /api/chat/end is a safe, idempotent follow-up ----------------
def test_chat_end_is_idempotent_after_automatic_save(client, monkeypatch):
    def fake_run_turn(session, text, debug=False):
        session.messages.append({"role": "user", "content": text})
        return "Goodbye!"

    def fake_finish(session):
        return conversation_end.finish_conversation(session, summarizer=lambda s: "A fake summary.")

    monkeypatch.setattr(main, "run_turn", fake_run_turn)
    monkeypatch.setattr(main, "finish_conversation", fake_finish)

    session_id = _start(client)
    client.post("/api/chat", json={"session_id": session_id, "message": "bye"})

    r1 = client.post("/api/chat/end", json={"session_id": session_id})
    assert r1.status_code == 200
    assert r1.json()["saved"] is True

    r2 = client.post("/api/chat/end", json={"session_id": session_id})
    assert r2.status_code == 200
    assert r2.json()["saved"] is True


# ---------------- the inactivity sweeper guarantees the save (reliability fix) ----------------
def test_sweep_saves_inactive_session_with_a_patient_message(client, monkeypatch):
    _fake_finish(monkeypatch)

    session_id = _start(client)
    session = sessions.get_session(session_id)
    session.messages.append({"role": "user", "content": "hello"})
    session.last_activity = datetime.now() - timedelta(minutes=11)

    saved = main.sweep_inactive_sessions(datetime.now())

    assert session_id in saved
    assert session.summary_saved is True

    wb = openpyxl.load_workbook(reports.REPORT_PATH)
    ws = wb[reports.SHEET]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert any(row[0] == session_id for row in rows)


def test_sweep_skips_inactive_session_without_a_patient_message(client, monkeypatch):
    _fake_finish(monkeypatch)

    session_id = _start(client)  # only the assistant greeting, no user message
    session = sessions.get_session(session_id)
    session.last_activity = datetime.now() - timedelta(minutes=11)

    saved = main.sweep_inactive_sessions(datetime.now())

    assert session_id not in saved
    assert session.summary_saved is False


def test_sweep_skips_a_recently_active_session(client, monkeypatch):
    _fake_finish(monkeypatch)

    session_id = _start(client)
    session = sessions.get_session(session_id)
    session.messages.append({"role": "user", "content": "hello"})
    session.last_activity = datetime.now()  # still active

    saved = main.sweep_inactive_sessions(datetime.now())

    assert session_id not in saved
    assert session.summary_saved is False


def test_sweep_skips_an_already_saved_session(client, monkeypatch):
    _fake_finish(monkeypatch)

    session_id = _start(client)
    session = sessions.get_session(session_id)
    session.messages.append({"role": "user", "content": "hello"})
    session.last_activity = datetime.now() - timedelta(minutes=11)
    session.summary_saved = True

    saved = main.sweep_inactive_sessions(datetime.now())

    assert session_id not in saved


# ---------------- two-tier timeout: booked sessions are saved sooner ----------------
def test_sweep_saves_booked_session_after_3_minutes_not_before(client, monkeypatch):
    _fake_finish(monkeypatch)

    session_id = _start(client)
    session = sessions.get_session(session_id)
    session.messages.append({"role": "user", "content": "please book me in"})
    session.booking = {"doctor": "Dr. Test", "speciality": "General Medicine"}
    session.last_activity = datetime.now() - timedelta(minutes=1)

    now = datetime.now()
    assert session_id not in main.sweep_inactive_sessions(now)  # only 1 minute idle

    saved = main.sweep_inactive_sessions(now + timedelta(minutes=3))  # now ~4 minutes idle
    assert session_id in saved
    assert session.summary_saved is True


def test_sweep_unbooked_session_waits_the_longer_timeout(client, monkeypatch):
    _fake_finish(monkeypatch)

    session_id = _start(client)
    session = sessions.get_session(session_id)
    session.messages.append({"role": "user", "content": "what are your opening hours?"})
    session.last_activity = datetime.now() - timedelta(minutes=4)

    now = datetime.now()
    # 4 minutes idle: past the booked timeout, but not the longer info-only one
    assert session_id not in main.sweep_inactive_sessions(now)
    assert session.summary_saved is False

    saved = main.sweep_inactive_sessions(now + timedelta(minutes=7))  # now ~11 minutes idle
    assert session_id in saved
    assert session.summary_saved is True


def test_sweep_retries_a_failed_save_on_the_next_sweep(client, monkeypatch):
    _fake_finish(monkeypatch)

    calls = {"n": 0}
    real_append = reports.append_conversation_row

    def flaky_append(row):
        calls["n"] += 1
        if calls["n"] == 1:
            return {"ok": False, "error": "Permission denied (Excel file open)."}
        return real_append(row)

    monkeypatch.setattr(reports, "append_conversation_row", flaky_append)

    session_id = _start(client)
    session = sessions.get_session(session_id)
    session.messages.append({"role": "user", "content": "hello"})
    session.last_activity = datetime.now() - timedelta(minutes=11)

    first_sweep = main.sweep_inactive_sessions(datetime.now())
    assert session_id not in first_sweep
    assert session.summary_saved is False

    second_sweep = main.sweep_inactive_sessions(datetime.now())
    assert session_id in second_sweep
    assert session.summary_saved is True


# ---------------- negative paths ----------------
def test_unknown_session_id_is_404(client):
    r = client.post("/api/chat", json={"session_id": "CONV-NOPE", "message": "hi"})
    assert r.status_code == 404

    r2 = client.post("/api/chat/end", json={"session_id": "CONV-NOPE"})
    assert r2.status_code == 404


def test_empty_or_whitespace_message_is_422(client):
    session_id = _start(client)
    r1 = client.post("/api/chat", json={"session_id": session_id, "message": ""})
    assert r1.status_code == 422

    r2 = client.post("/api/chat", json={"session_id": session_id, "message": "   "})
    assert r2.status_code == 422


def test_invoice_filename_safety(client):
    r1 = client.get("/invoices/..%2Fsecret.pdf")
    assert r1.status_code in (400, 404)

    r2 = client.get("/invoices/not-a-pdf.txt")
    assert r2.status_code == 400

    r3 = client.get("/invoices/does-not-exist.pdf")
    assert r3.status_code == 404


def test_health_and_hospital(client):
    r1 = client.get("/api/health")
    assert r1.status_code == 200
    assert r1.json()["status"] == "ok"
    assert r1.json()["model"]

    r2 = client.get("/api/hospital")
    assert r2.status_code == 200
    body = r2.json()
    assert body["name"] and body["phone"] and body["emergency_number"]
    assert body["address"]
    assert "Cardiology" in (body["departments"] or "")
    assert "Pharmacy" in (body["facilities"] or "")


# ---------------- regression: prompts.py must read the real hospital_info row ----------------
def test_chat_start_greeting_uses_real_hospital_name(client):
    r = client.post("/api/chat/start")
    assert r.status_code == 200
    assert "Lotus Care Hospital" in r.json()["greeting"]
